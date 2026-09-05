#!/usr/bin/env python3
"""
build_stabilization_client_summary.py — rebuilds the "Stabilization - Summary"
sheet in the client-facing IIRM_Stabilization_Summary_<date>.xlsx file from its
own "Stabilization - Items" sheet.

This is a standalone Stabilization item list (not filtered out of the main SOW
tracker by Area), so it reuses build_iirm_summary's Config/Helper/Stablization
Summary builders pointed at this file's own source sheet instead of the main
tracker's "SOW, Feedbacks-Bug & CR_Items".

Usage:
  python3 build_stabilization_client_summary.py --file IIRM_Stabilization_Summary_07-Aug-2026.xlsx
"""

import argparse
import warnings
from pathlib import Path

import openpyxl
import pandas as pd

import build_iirm_summary as bis

warnings.filterwarnings("ignore")

ITEMS_SHEET = "Stabilization - Items"
SUMMARY_SHEET = "Stabilization - Summary"


def main(file_path):
    wb = openpyxl.load_workbook(file_path)
    if ITEMS_SHEET not in wb.sheetnames:
        raise SystemExit(f"ERROR: '{ITEMS_SHEET}' sheet not found in {file_path}")

    df_src = pd.read_excel(file_path, sheet_name=ITEMS_SHEET)
    df_src["_bucket"] = df_src["Status"].apply(
        lambda s: bis.STATUS_MAP.get(str(s).strip(), ("Unknown", "N/A"))[1]
        if pd.notna(s) else "N/A"
    )
    print(f"Rows in '{ITEMS_SHEET}': {len(df_src)}")

    ws_items = wb[ITEMS_SHEET]
    last_row = ws_items.max_row
    while last_row > 1 and all(
        ws_items.cell(last_row, c).value in (None, "") for c in range(1, 8)
    ):
        last_row -= 1
    print(f"Last data row: {last_row}")

    req_date_col = None
    for c in ws_items[1]:
        if c.value and str(c.value).strip().lower() in (
                "requested date", "request date", "requested", "req date"):
            req_date_col = c.column_letter
            break
    print(f"Requested Date column: {req_date_col or 'not found'}")

    # Strip any prior-run sheets — keep the source untouched, matching the
    # main script's rebuild-from-scratch design (no accumulated Helper1/Config1).
    for name in (SUMMARY_SHEET, "Stablization - Summary", "Helper", "Config"):
        if name in wb.sheetnames:
            del wb[name]

    bis.build_config(wb)
    bis.build_helper(wb, ITEMS_SHEET, last_row, req_date_col)

    stab_mask = df_src["Area"].astype(str).str.strip() == bis.STAB_AREA_LABEL
    present_features = sorted(
        df_src.loc[stab_mask, "Feature"].dropna().astype(str).str.strip().unique()
    )
    print(f"Features found: {present_features}")

    bis.build_stablization_summary(
        wb, df_src, present_features,
        sheet_name=SUMMARY_SHEET, source_sheet=ITEMS_SHEET,
    )

    desired_order = [ITEMS_SHEET, SUMMARY_SHEET, "Helper", "Config"]
    wb._sheets.sort(
        key=lambda s: desired_order.index(s.title) if s.title in desired_order else 99
    )

    wb.save(file_path)
    print(f"\n✅  Saved -> {file_path}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--file", required=True, type=Path, help="Path to the client Stabilization file")
    args = parser.parse_args()
    main(args.file)
