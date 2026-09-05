#!/usr/bin/env python3
"""
jira_report_extract.py — on-demand extract of Jira issues into the
"Jira Report" sheet of the IIRM SOW tracker Excel.

See Jira-Report-Extract.md (same folder) for the column definitions,
lookup values, and full usage notes.

Usage:
  python jira_report_extract.py --jql "project = IIRM AND ..."          # incremental
  python jira_report_extract.py --jql "project = IIRM AND ..." --full   # full pull
  python jira_report_extract.py --jql "..." --file /path/to/other.xlsx  # override target

Requires JIRA_SERVER, JIRA_EMAIL, JIRA_TOKEN in the environment — direnv
loads these automatically from .env.jira when a terminal is opened in the
repo root (see Jira-Sync-Setup.md).
"""

import argparse
import os
import re
import sys
import time
from datetime import datetime, timedelta
from pathlib import Path

import openpyxl
import requests

SCRIPT_DIR = Path(__file__).parent
DEFAULT_EXCEL = SCRIPT_DIR / "IIRM_SOW_Executive_Summary_07-Aug-2026.xlsx"
FROZEN_FILE_MARKER = "24-Jul-2026"
SHEET_NAME = "Jira Report"
LOG_PATH = SCRIPT_DIR / "Jira-Report-Extract.md"
LAST_EXTRACT_HEADING = "## Last extract"
ISO_FORMAT = "%Y-%m-%dT%H:%M:%S%z"
JQL_DATE_FORMAT = "%Y-%m-%d %H:%M"
SAFETY_BUFFER_MINUTES = 15  # subtract from last-extract time before querying, to cover clock/timezone skew
REQUIRED_COLUMNS = ["sl. #", "jira id", "issue type", "summary", "status", "reporter", "product", "parent"]
SEARCH_FIELDS = "summary,issuetype,status,reporter,customfield_10228,parent"
PAGE_SIZE = 100


class Client:
    """Thin wrapper around /rest/api/3/search/jql — Jira Cloud retired the
    older offset-paginated /rest/api/2/search (410 Gone), so pagination here
    is cursor-based (nextPageToken), not startAt-based."""

    def __init__(self, server: str, email: str, token: str):
        self.base = server.rstrip("/")
        self.auth = (email, token)

    def search(self, jql: str):
        page_token = None
        while True:
            params = {"jql": jql, "maxResults": PAGE_SIZE, "fields": SEARCH_FIELDS}
            if page_token:
                params["nextPageToken"] = page_token
            for attempt in range(5):
                r = requests.get(f"{self.base}/rest/api/3/search/jql", params=params, auth=self.auth)
                if r.status_code == 429:
                    time.sleep(2 ** attempt)
                    continue
                break
            if r.status_code != 200:
                sys.exit(f"ERROR: Jira search failed ({r.status_code}): {r.text[:300]}")
            data = r.json()
            batch = data.get("issues", [])
            yield from batch
            if data.get("isLast") or not batch:
                return
            page_token = data.get("nextPageToken")


def make_client() -> Client:
    for name in ("JIRA_SERVER", "JIRA_EMAIL", "JIRA_TOKEN"):
        if not os.environ.get(name):
            sys.exit(
                f"ERROR: {name} not set. Open a terminal in the repo root "
                "(direnv loads .env.jira automatically) or export it manually."
            )
    server = os.environ["JIRA_SERVER"]
    if not server.startswith("https://"):
        sys.exit("ERROR: JIRA_SERVER must use HTTPS.")
    client = Client(server, os.environ["JIRA_EMAIL"], os.environ["JIRA_TOKEN"])
    r = requests.get(f"{client.base}/rest/api/3/myself", auth=client.auth)
    if r.status_code != 200:
        sys.exit(
            f"ERROR: Jira connection failed ({r.status_code}): {r.text[:300]}\n"
            "If this is a 401: the token in .env.jira may be the one flagged as burned "
            "in Jira-Sync-Setup.md. Generate a fresh token at "
            "https://id.atlassian.com/manage-profile/security/api-tokens, update .env.jira "
            "directly (never paste it into chat), then revoke the old one."
        )
    return client


def read_last_extract() -> datetime | None:
    if not LOG_PATH.exists():
        return None
    text = LOG_PATH.read_text()
    m = re.search(rf"{re.escape(LAST_EXTRACT_HEADING)}\s*\n\s*(.+)", text)
    if not m:
        return None
    raw = m.group(1).strip()
    try:
        return datetime.strptime(raw, ISO_FORMAT)
    except ValueError:
        return None  # e.g. the initial "Never (script has not been run yet)" placeholder


def write_last_extract(timestamp: datetime) -> None:
    stamp = timestamp.strftime(ISO_FORMAT)
    if not LOG_PATH.exists():
        sys.exit(f"ERROR: {LOG_PATH} not found — cannot record extract timestamp.")
    text = LOG_PATH.read_text()
    if LAST_EXTRACT_HEADING in text:
        text = re.sub(
            rf"{re.escape(LAST_EXTRACT_HEADING)}\s*\n\s*.+",
            f"{LAST_EXTRACT_HEADING}\n\n{stamp}",
            text,
            count=1,
        )
    else:
        text = text.rstrip() + f"\n\n{LAST_EXTRACT_HEADING}\n\n{stamp}\n"
    LOG_PATH.write_text(text)


def fetch_issues(client: Client, jql: str) -> list[dict]:
    issues = []
    for issue in client.search(jql):
        issues.append(issue)
        if len(issues) % 500 == 0:
            print(f"  ...fetched {len(issues)} issue(s) so far")
    return issues


def issue_to_row(issue: dict) -> dict:
    f = issue["fields"]
    issuetype = f.get("issuetype") or {}
    status = f.get("status") or {}
    reporter = f.get("reporter") or {}
    product = f.get("customfield_10228") or {}
    parent = f.get("parent")
    parent_str = f"{parent['key']} — {parent['fields']['summary']}" if parent else ""
    return {
        "jira_id": issue["key"],
        "issue_type": issuetype.get("name", ""),
        "summary": f.get("summary") or "",
        "status": status.get("name", ""),
        "reporter": reporter.get("displayName", ""),
        "product": product.get("value", ""),
        "parent": parent_str,
    }


def update_sheet(excel_path: Path, rows: list[dict]) -> tuple[int, int]:
    wb = openpyxl.load_workbook(excel_path)
    if SHEET_NAME not in wb.sheetnames:
        sys.exit(f"ERROR: sheet '{SHEET_NAME}' not found in {excel_path}")
    ws = wb[SHEET_NAME]
    hdr = {str(c.value).strip().lower(): c.column for c in ws[1] if c.value}
    missing = [h for h in REQUIRED_COLUMNS if h not in hdr]
    if missing:
        sys.exit(f"ERROR: sheet '{SHEET_NAME}' is missing expected column(s): {missing}")

    existing_row_by_id = {}
    max_sl = 0
    max_row = ws.max_row
    for r in range(2, max_row + 1):
        jid = ws.cell(r, hdr["jira id"]).value
        if jid:
            existing_row_by_id[jid] = r
        sl = ws.cell(r, hdr["sl. #"]).value
        if isinstance(sl, (int, float)):
            max_sl = max(max_sl, int(sl))

    added = 0
    refreshed = 0
    next_row = max_row + 1
    for row in rows:
        jid = row["jira_id"]
        if jid in existing_row_by_id:
            r = existing_row_by_id[jid]
            refreshed += 1
        else:
            r = next_row
            next_row += 1
            max_sl += 1
            ws.cell(r, hdr["sl. #"], max_sl)
            ws.cell(r, hdr["jira id"], jid)
            added += 1
        ws.cell(r, hdr["issue type"], row["issue_type"])
        ws.cell(r, hdr["summary"], row["summary"])
        ws.cell(r, hdr["status"], row["status"])
        ws.cell(r, hdr["reporter"], row["reporter"])
        ws.cell(r, hdr["product"], row["product"])
        ws.cell(r, hdr["parent"], row["parent"])

    wb.save(excel_path)
    return added, refreshed


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--jql", required=True, help='Base JQL filter, e.g. "project = IIRM AND labels = X"')
    parser.add_argument("--full", action="store_true", help="Ignore the last-extract timestamp; pull everything matching --jql")
    parser.add_argument("--file", type=Path, default=DEFAULT_EXCEL, help="Target Excel file (defaults to the current working file)")
    args = parser.parse_args()

    if FROZEN_FILE_MARKER in str(args.file):
        sys.exit(f"ERROR: refusing to write to a frozen file ({args.file}). This file must never be modified.")

    run_started = datetime.now().astimezone()
    jql = args.jql
    if not args.full:
        last = read_last_extract()
        if last:
            cutoff = last - timedelta(minutes=SAFETY_BUFFER_MINUTES)
            jql = f'({jql}) AND updated >= "{cutoff.strftime(JQL_DATE_FORMAT)}"'
            print(f"Incremental mode: pulling issues updated since {cutoff.strftime(JQL_DATE_FORMAT)} "
                  f"(last extract was {last.strftime(ISO_FORMAT)}, minus a {SAFETY_BUFFER_MINUTES}-min safety buffer)")
        else:
            print("No valid prior extract timestamp found — pulling full scope for --jql.")
    else:
        print("Full mode: ignoring last-extract timestamp.")

    if "order by" not in jql.lower():
        jql = f"{jql} ORDER BY key ASC"

    print(f"JQL: {jql}")
    client = make_client()
    issues = fetch_issues(client, jql)
    print(f"Fetched {len(issues)} issue(s) from Jira.")

    if issues:
        rows = [issue_to_row(i) for i in issues]
        added, refreshed = update_sheet(args.file, rows)
        print(f"Sheet updated: {added} new row(s) added, {refreshed} existing row(s) refreshed.")
    else:
        print("Nothing to update.")

    write_last_extract(run_started)
    print(f"Last-extract timestamp recorded: {run_started.strftime(ISO_FORMAT)}")


if __name__ == "__main__":
    main()
